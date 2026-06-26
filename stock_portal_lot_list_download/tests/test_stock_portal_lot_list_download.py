# Copyright 2025 Tecnativa - Pilar Vargas
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl).
import io

from openpyxl import load_workbook

from odoo.tests import Form, HttpCase, tagged

from odoo.addons.mail.tests.common import mail_new_test_user


@tagged("-at_install", "post_install")
class TestPortalPickingXlsx(HttpCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env = cls.env(context=dict(cls.env.context, mail_notrack=True))
        cls.portal_user = mail_new_test_user(
            cls.env, login="portal_user", groups="base.group_portal"
        )
        cls.product = cls.env["product.product"].create(
            {
                "name": "Product with lot",
                "is_storable": True,
                "tracking": "lot",
                "list_price": 50.0,
            }
        )
        cls.product_no_lot = cls.env["product.product"].create(
            {
                "name": "Product without lot",
                "is_storable": True,
                "tracking": "none",
                "list_price": 30.0,
            }
        )
        cls.lot = cls.env["stock.lot"].create(
            {
                "name": "LOT001",
                "product_id": cls.product.id,
            }
        )
        sale_form = Form(cls.env["sale.order"])
        sale_form.partner_id = cls.portal_user.partner_id
        with sale_form.order_line.new() as line:
            line.product_id = cls.product
            line.product_uom_qty = 2
        with sale_form.order_line.new() as line:
            line.product_id = cls.product_no_lot
            line.product_uom_qty = 1
        cls.sale_order = sale_form.save()
        cls.sale_order.action_confirm()
        cls.picking = cls.sale_order.picking_ids[:1]
        for move in cls.picking.move_ids:
            move.quantity = move.product_uom_qty
        tracked_line = cls.picking.move_line_ids.filtered(
            lambda line: line.product_id.tracking != "none"
        )
        tracked_line.lot_id = cls.lot.id
        cls.picking.button_validate()
        no_lot_sale_form = Form(cls.env["sale.order"])
        no_lot_sale_form.partner_id = cls.portal_user.partner_id
        with no_lot_sale_form.order_line.new() as line:
            line.product_id = cls.product_no_lot
            line.product_uom_qty = 1
        cls.no_lot_sale_order = no_lot_sale_form.save()
        cls.no_lot_sale_order.action_confirm()
        cls.no_lot_picking = cls.no_lot_sale_order.picking_ids[:1]
        for move in cls.no_lot_picking.move_ids:
            move.quantity = move.product_uom_qty
        cls.no_lot_picking.button_validate()

    def test_portal_download_xlsx_file(self):
        # Ensure that the Excel file can be downloaded from the portal
        self.authenticate("portal_user", "portal_user")
        url = f"/my/picking/xlsx/{self.picking.id}"
        response = self.url_open(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.headers.get("Content-Type", ""),
        )
        self.assertTrue(response.content)
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        # Verify headers and content data
        self.assertEqual(sheet.cell(row=3, column=1).value, "Product")
        self.assertEqual(sheet.cell(row=3, column=2).value, "Lot/Serial Number")
        self.assertEqual(sheet.cell(row=3, column=3).value, "Quantity")
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.cell(row=4, column=1).value, self.product.display_name)
        self.assertEqual(sheet.cell(row=4, column=2).value, self.lot.name)
        self.assertEqual(sheet.cell(row=4, column=3).value, 2.0)
        all_product_names = [sheet.cell(row=row, column=1).value for row in range(4, 5)]
        self.assertNotIn(self.product_no_lot.display_name, all_product_names)

    def test_public_download_xlsx_file_with_access_token(self):
        self.authenticate(None, None)
        self.sale_order._portal_ensure_token()
        url = (
            f"/my/picking/xlsx/{self.picking.id}"
            f"?access_token={self.sale_order.access_token}"
        )
        response = self.url_open(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.headers.get("Content-Type", ""),
        )

    def test_portal_sale_order_link_visibility(self):
        self.authenticate("portal_user", "portal_user")
        response = self.url_open(f"/my/orders/{self.sale_order.id}")
        self.assertIn(f"/my/picking/xlsx/{self.picking.id}", response.text)
        response = self.url_open(f"/my/orders/{self.no_lot_sale_order.id}")
        self.assertNotIn(f"/my/picking/xlsx/{self.no_lot_picking.id}", response.text)

    def test_portal_download_xlsx_requires_picking_access(self):
        self.authenticate(None, None)
        response = self.url_open(
            f"/my/picking/xlsx/{self.picking.id}", allow_redirects=False
        )
        self.assertEqual(response.status_code, 404)
